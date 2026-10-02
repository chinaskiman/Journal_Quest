"""Excel exporters.

Three flavours:
  * journal  - faithful reproduction of the Journal sheet, including live
               formulas for every calculated column (the risk-based Net
               PnL model, and corrected streak formulas).
  * analytics- reproduction of the Analytics sheet with values computed
               by the app.
  * backup   - pure values (Journal + Options sheets) suitable for
               round-tripping through the importer.
"""
from __future__ import annotations

import io
from datetime import timedelta
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy.orm import Session

from .models import OptionValue, Trade
from .serializers import serialize_trades

HEADER_FILL = PatternFill("solid", fgColor="1F2933")
HEADER_FONT = Font(color="FFFFFF", bold=True)
SECTION_FILL = PatternFill("solid", fgColor="233A50")
SECTION_FONT = Font(color="FFFFFF", bold=True, size=12)
GROUP_HEADER_FILL = PatternFill("solid", fgColor="2E4A66")
GROUP_HEADER_FONT = Font(color="FFFFFF", bold=True)

# (header, width, number_format)
JOURNAL_LAYOUT: list[tuple[str, int, str | None]] = [
    ("#", 6, "0"),
    ("Symbol", 12, None),
    ("Opening Day", 22, "[$-F800]dddd\\,\\ mmmm\\ dd\\,\\ yyyy"),
    ("Opening Time", 12, "h:mm;@"),
    ("Type", 18, None),
    ("Direction", 10, None),
    ("Volume", 9, None),
    ("Entry Price", 11, "0.00000"),
    ("SL Price", 10, "0.00000"),
    ("SL ATR Multiplier", 13, "0.0"),
    ("TP Price", 10, "0.00000"),
    ("R:R", 7, "0.0"),
    ("Risk", 9, '"$"#,##0.00'),
    ("4H Bias", 9, None),
    ("1H Bias", 9, None),
    ("Story Behind Candles", 34, None),
    ("Opening RSI Condition", 28, None),
    ("ATR", 9, "0.00000"),
    ("Mental State", 13, None),
    ("Why I Entered", 34, None),
    ("Closing Day", 22, "[$-F800]dddd\\,\\ mmmm\\ dd\\,\\ yyyy"),
    ("Closing Time", 12, "h:mm;@"),
    ("Closing Price", 12, "0.00000"),
    ("Closing RSI", 11, "0.00"),
    ("Why I Closed", 34, None),
    ("Final R:R", 10, "0.0"),
    ("Commission", 11, '"$"#,##0.00'),
    ("Swap", 8, '"$"#,##0.00'),
    ("Net PnL", 11, '"$"#,##0.00'),
    ("Status", 9, None),
    ("Duration", 12, None),
    ("Note", 34, None),
    ("Win Streak (helper)", 12, "0"),
    ("Loss Streak (helper)", 12, "0"),
]

#: trade field (or None for formula columns) per Journal column index.
JOURNAL_VALUE_FIELDS: list[str | None] = [
    "seq",
    "symbol",
    "opening_day",
    "opening_time",
    "setup_type",
    "direction",
    "volume",
    "entry_price",
    "sl_price",
    None,  # J  SL ATR Multiplier (formula)
    "tp_price",
    None,  # L  R:R (formula)
    "risk",
    "bias_4h",
    "bias_1h",
    "story_behind_candles",
    "opening_rsi_condition",
    "atr",
    "mental_state",
    "why_i_entered",
    "closing_day",
    "closing_time",
    "closing_price",
    "closing_rsi",
    "why_i_closed",
    None,  # Z  Final R:R (formula)
    "commission",
    "swap",
    None,  # AC Net PnL (formula)
    None,  # AD Status (formula)
    None,  # AE Duration (formula)
    "note",
    None,  # AG Win streak (formula)
    None,  # AH Loss streak (formula)
]


def _calculated_formulas(excel_row: int) -> dict[int, str]:
    """Live Excel formulas for the calculated columns, keyed by column
    index (1-based, matching JOURNAL_LAYOUT)."""
    r = excel_row
    return {
        10: f'=IF(OR(H{r}="",I{r}="",R{r}="",R{r}=0),"",ABS(H{r}-I{r})/R{r})',
        12: (
            f'=IF(OR(H{r}="",I{r}="",K{r}=""),"",IF(F{r}="Short",'
            f'IF((I{r}-H{r})<=0,"",(H{r}-K{r})/(I{r}-H{r})),'
            f'IF((H{r}-I{r})<=0,"",(K{r}-H{r})/(H{r}-I{r}))))'
        ),
        26: (
            f'=IF(OR(H{r}="",I{r}="",W{r}=""),"",IF(F{r}="Short",'
            f'IF((I{r}-H{r})<=0,"",(H{r}-W{r})/(I{r}-H{r})),'
            f'IF((H{r}-I{r})<=0,"",(W{r}-H{r})/(H{r}-I{r}))))'
        ),
        # Net PnL: risk-based model (Final R:R x Risk - Commission + Swap)
        29: (
            f'=IF(OR(Z{r}="",M{r}=""),"",Z{r}*M{r}'
            f'-IF(AA{r}="",0,AA{r})+IF(AB{r}="",0,AB{r}))'
        ),
        30: (
            f'=IF(B{r}="","",IF(W{r}="","Open",IF(AC{r}="","Open",'
            f'IF(AC{r}>0,"Win",IF(AC{r}<0,"Loss","BE")))))'
        ),
        31: (
            f'=IF(OR(C{r}="",D{r}="",U{r}="",V{r}=""),"",'
            f"(U{r}+V{r})-(C{r}+D{r}))"
        ),
        33: f'=IF(AD{r}="Win",IF(ISNUMBER(AG{r-1}),AG{r-1}+1,1),0)',
        34: f'=IF(AD{r}="Loss",IF(ISNUMBER(AH{r-1}),AH{r-1}+1,1),0)',
    }


def _write_journal_sheet(ws, rows: list[dict[str, Any]], formulas: bool) -> None:
    # header
    for col_idx, (header, width, _fmt) in enumerate(JOURNAL_LAYOUT, start=1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")
        ws.column_dimensions[get_column_letter(col_idx)].width = width
    ws.row_dimensions[1].height = 24

    for i, row in enumerate(rows):
        excel_row = i + 2
        calc = _calculated_formulas(excel_row) if formulas else {}
        for col_idx, field in enumerate(JOURNAL_VALUE_FIELDS, start=1):
            value: Any
            if field is None:
                if not formulas:
                    derived_key = {
                        10: "sl_atr_multiplier",
                        12: "planned_rr",
                        26: "final_rr",
                        29: "net_pnl",
                        30: "status",
                        31: "duration_seconds",
                        33: "win_streak",
                        34: "loss_streak",
                    }[col_idx]
                    value = row.get(derived_key)
                else:
                    value = calc.get(col_idx)
            else:
                value = row.get(field)
                if field == "duration_seconds" and value is not None:
                    # express as a native Excel duration
                    value = timedelta(seconds=value)
            cell = ws.cell(row=excel_row, column=col_idx, value=value)
            fmt = JOURNAL_LAYOUT[col_idx - 1][2]
            if isinstance(value, timedelta):
                cell.number_format = "[h]:mm:ss"
            elif fmt and value is not None:
                cell.number_format = fmt

    ws.freeze_panes = "C2"


def _get_rows(db: Session) -> list[dict[str, Any]]:
    trades = db.query(Trade).order_by(Trade.seq, Trade.id).all()
    return serialize_trades(trades)


def export_journal_xlsx(db: Session) -> bytes:
    """Journal sheet with live formulas for all calculated columns."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Journal"
    _write_journal_sheet(ws, _get_rows(db), formulas=True)
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def export_backup_xlsx(db: Session) -> bytes:
    """Full backup: Journal (pure values) + Options sheets - can be
    re-imported through the import endpoint."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Journal"
    _write_journal_sheet(ws, _get_rows(db), formulas=False)

    opt = wb.create_sheet("Options")
    from .importer import OPTIONS_HEADER_TO_KEY

    headers = {v: k for k, v in OPTIONS_HEADER_TO_KEY.items()}
    for idx, (list_key, header) in enumerate(headers.items(), start=1):
        cell = opt.cell(row=2, column=idx, value=header)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        opt.column_dimensions[get_column_letter(idx)].width = 22
        for row_idx, (value,) in enumerate(
            db.query(OptionValue.value)
            .filter(OptionValue.list_key == list_key)
            .order_by(OptionValue.position, OptionValue.id),
            start=3,
        ):
            opt.cell(row=row_idx, column=idx, value=value)

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


# ---------------------------------------------------------------------------
# Analytics sheet reproduction
# ---------------------------------------------------------------------------
_MONEY = '"$"#,##0.00'
_RATIO = "0.00"
_PERCENT = "0.0%"


def _option_lists(db: Session) -> dict[str, list[str]]:
    lists: dict[str, list[str]] = {}
    rows = (
        db.query(OptionValue)
        .order_by(OptionValue.list_key, OptionValue.position, OptionValue.id)
        .all()
    )
    for row in rows:
        lists.setdefault(row.list_key, []).append(row.value)
    return lists


def _label(ws, row: int, label: str, value, fmt: str | None = None) -> None:
    ws.cell(row=row, column=1, value=label)
    cell = ws.cell(row=row, column=2, value=value)
    if fmt and value is not None:
        cell.number_format = fmt


def _section(ws, row: int, title: str) -> None:
    for col in range(1, 7):
        cell = ws.cell(row=row, column=col)
        cell.fill = SECTION_FILL
    cell = ws.cell(row=row, column=1, value=title)
    cell.font = SECTION_FONT


def _group_table(ws, row: int, title: str, groups: list[dict[str, Any]]) -> int:
    """Write one performance table; returns the next free row."""
    _section(ws, row, title)
    headers = ["", "Total Trades", "Closed", "Win Rate", "Avg Realized R", "Net P&L"]
    for col, header in enumerate(headers, start=1):
        cell = ws.cell(row=row + 1, column=col, value=header or None)
        cell.fill = GROUP_HEADER_FILL
        cell.font = GROUP_HEADER_FONT
    for offset, group in enumerate(groups):
        r = row + 2 + offset
        ws.cell(row=r, column=1, value=group["group"])
        ws.cell(row=r, column=2, value=group["total_trades"])
        ws.cell(row=r, column=3, value=group["closed"])
        if group["win_rate"] is not None:
            cell = ws.cell(row=r, column=4, value=group["win_rate"])
            cell.number_format = _PERCENT
        if group["avg_realized_r"] is not None:
            cell = ws.cell(row=r, column=5, value=group["avg_realized_r"])
            cell.number_format = _RATIO
        cell = ws.cell(row=r, column=6, value=group["net_pnl"])
        cell.number_format = _MONEY
    return row + 2 + len(groups) + 1


def export_analytics_xlsx(db: Session, payload: dict[str, Any]) -> bytes:
    """Reproduce the Analytics sheet layout with live computed values."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Analytics"

    title = ws.cell(row=1, column=1, value="Performance Analytics")
    title.font = Font(bold=True, size=14)
    subtitle = ws.cell(
        row=2,
        column=1,
        value=(
            "All metrics are calculated automatically from the journal and "
            "update whenever trades change."
        ),
    )
    subtitle.font = Font(italic=True, color="808080")

    overview = payload["overview"]
    _section(ws, 4, "Overview")
    _label(ws, 5, "Total trades logged", overview["total_trades"])
    _label(ws, 6, "Closed trades", overview["closed_trades"])
    _label(ws, 7, "Open trades", overview["open_trades"])
    _label(ws, 8, "Wins", overview["wins"])
    _label(ws, 9, "Losses", overview["losses"])
    _label(ws, 10, "Break-even (BE) trades", overview["break_even"])
    _label(ws, 11, "Win rate", overview["win_rate"], _PERCENT)

    pnl = payload["pnl"]
    _section(ws, 13, "Profit & Loss")
    _label(ws, 14, "Total Net P&L", pnl["total_net_pnl"], _MONEY)
    _label(ws, 15, "Gross profit (wins only)", pnl["gross_profit"], _MONEY)
    _label(ws, 16, "Gross loss (losses only)", pnl["gross_loss"], _MONEY)
    _label(ws, 17, "Profit factor", pnl["profit_factor"], _RATIO)
    _label(ws, 18, "Average win", pnl["average_win"], _MONEY)
    _label(ws, 19, "Average loss", pnl["average_loss"], _MONEY)
    _label(ws, 20, "Expectancy per closed trade", pnl["expectancy"], _MONEY)
    _label(ws, 21, "Best trade (max net P&L)", pnl["best_trade"], _MONEY)
    _label(ws, 22, "Worst trade (min net P&L)", pnl["worst_trade"], _MONEY)

    risk = payload["risk"]
    _section(ws, 24, "Risk & Ratios")
    _label(ws, 25, "Average planned R:R", risk["average_planned_rr"], _RATIO)
    _label(ws, 26, "Average realized R", risk["average_realized_r"], _RATIO)
    _label(ws, 27, "Average risk per trade", risk["average_risk"], _MONEY)
    _label(ws, 28, "Max consecutive winning trades", risk["max_consecutive_wins"])
    _label(ws, 29, "Max consecutive losing trades", risk["max_consecutive_losses"])
    _label(ws, 30, "Total commissions", risk["total_commission"], _MONEY)
    _label(ws, 31, "Total swap", risk["total_swap"], _MONEY)

    cursor = _group_table(ws, 36, "Performance by Symbol", payload["by_symbol"])
    cursor = _group_table(ws, cursor + 1, "Performance by Direction", payload["by_direction"])
    cursor = _group_table(ws, cursor + 1, "Performance by Setup Type", payload["by_setup_type"])
    cursor = _group_table(ws, cursor + 1, "Performance by Mental State", payload["by_mental_state"])

    note = ws.cell(
        row=cursor + 1,
        column=1,
        value="Calculated live by TradeQuest - not stored in the database.",
    )
    note.font = Font(italic=True, color="808080")

    widths = {1: 34, 2: 16, 3: 12, 4: 12, 5: 16, 6: 14}
    for col, width in widths.items():
        ws.column_dimensions[get_column_letter(col)].width = width

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()

