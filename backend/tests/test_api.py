"""REST API tests: CRUD, validation, filtering, sorting, pagination,
options management, import/export, and the live-analytics contract."""
from __future__ import annotations

TRADE = {
    "symbol": "EURUSD",
    "direction": "Long",
    "opening_day": "2026-09-20",
    "opening_time": "09:30",
    "setup_type": "Breakout",
    "volume": 1.0,
    "entry_price": 100.0,
    "sl_price": 95.0,
    "tp_price": 115.0,
    "risk": 50.0,
    "closing_day": "2026-09-21",
    "closing_time": "10:30",
    "closing_price": 110.0,
    "commission": 1.5,
    "swap": 0.25,
}


def _seed(client):
    for list_key, values in (
        ("symbol", ["EURUSD", "XAUUSD"]),
        ("setup_type", ["Breakout", "Rejection"]),
        ("bias", ["Bullish", "Bearish", "Neutral"]),
        ("rsi_condition", ["Rising from Support", "Exiting Oversold"]),
        ("mental_state", ["Calm", "Confident"]),
    ):
        for value in values:
            response = client.post(f"/api/options/{list_key}", json={"value": value})
            assert response.status_code == 201, response.text


def _post_trade(client, **overrides) -> dict:
    payload = dict(TRADE, **overrides)
    response = client.post("/api/trades", json=payload)
    assert response.status_code == 201, response.text
    return response.json()


# --------------------------------------------------------------------- health
def test_health(client):
    response = client.get("/api/health")
    assert response.status_code == 200
    assert response.json()["status"] == "ok"


# -------------------------------------------------------------------- create
def test_create_and_read_trade(client):
    _seed(client)
    body = _post_trade(client)
    assert body["seq"] == 1
    assert body["status"] == "Win"
    assert abs(body["planned_rr"] - 3.0) < 1e-9
    assert abs(body["final_rr"] - 2.0) < 1e-9
    assert abs(body["net_pnl"] - 98.75) < 1e-9
    assert body["duration_seconds"] == 90000  # 1d 1h
    assert body["duration_display"] == "1d 1h"

    got = client.get(f"/api/trades/{body['id']}")
    assert got.status_code == 200
    assert got.json()["net_pnl"] == body["net_pnl"]


def test_create_open_trade(client):
    _seed(client)
    payload = {k: v for k, v in TRADE.items() if not k.startswith("closing")}
    response = client.post("/api/trades", json=payload)
    assert response.status_code == 201, response.text
    body = response.json()
    assert body["status"] == "Open"
    assert body["net_pnl"] is None
    assert body["final_rr"] is None


def test_create_rejects_unknown_symbol(client):
    _seed(client)
    response = client.post("/api/trades", json=dict(TRADE, symbol="NOT_A_SYMBOL"))
    assert response.status_code == 422
    assert "not in the Options" in response.json()["detail"]


def test_create_rejects_bad_direction(client):
    _seed(client)
    response = client.post("/api/trades", json=dict(TRADE, direction="Up"))
    assert response.status_code == 422


def test_create_rejects_unlisted_rsi_condition(client):
    _seed(client)
    response = client.post(
        "/api/trades", json=dict(TRADE, opening_rsi_condition="Not An Option")
    )
    assert response.status_code == 422


def test_create_accepts_multiple_rsi_conditions(client):
    _seed(client)
    body = _post_trade(
        client, opening_rsi_condition="Rising from Support\nExiting Oversold"
    )
    assert body["opening_rsi_condition"].count("\n") == 1


# -------------------------------------------------------------------- update
def test_patch_updates_derived_fields(client):
    _seed(client)
    body = _post_trade(client)
    response = client.patch(f"/api/trades/{body['id']}", json={"closing_price": 95.0})
    assert response.status_code == 200
    updated = response.json()
    assert updated["final_rr"] == -1.0
    assert updated["status"] == "Loss"
    assert abs(updated["net_pnl"] - (-51.25)) < 1e-9  # -1*50 - 1.5 + 0.25


def test_patch_can_keep_legacy_unlisted_rsi_values(client):
    """Imported trades may keep hand-typed RSI variants when edited."""
    _seed(client)
    from app.db import SessionLocal
    from app.models import Trade

    session = SessionLocal()
    trade = Trade(
        seq=1,
        symbol="EURUSD",
        direction="Long",
        opening_rsi_condition="Rising From Support\nExiting oversold",
    )
    session.add(trade)
    session.commit()
    trade_id = trade.id
    session.close()

    response = client.patch(
        f"/api/trades/{trade_id}", json={"note": "edited, RSI untouched"}
    )
    assert response.status_code == 200
    assert "Rising From Support" in response.json()["opening_rsi_condition"]


def test_delete_renumbers_seq(client):
    _seed(client)
    ids = [_post_trade(client)["id"] for _ in range(3)]
    assert client.delete(f"/api/trades/{ids[0]}").status_code == 204
    rows = client.get("/api/trades").json()["rows"]
    assert [r["seq"] for r in rows] == [1, 2]


def test_duplicate_trade(client):
    _seed(client)
    body = _post_trade(client)
    response = client.post(f"/api/trades/{body['id']}/duplicate")
    assert response.status_code == 201
    clone = response.json()
    assert clone["id"] != body["id"]
    assert clone["seq"] == 2
    assert clone["symbol"] == body["symbol"]


def test_duplicate_open_trade_has_no_pnl(client):
    _seed(client)
    payload = {k: v for k, v in TRADE.items() if not k.startswith("closing")}
    body = client.post("/api/trades", json=payload).json()
    clone = client.post(f"/api/trades/{body['id']}/duplicate").json()
    assert clone["status"] == "Open"
    assert clone["net_pnl"] is None


# -------------------------------------------------------------------- reorder
def test_reorder_recomputes_streaks(client):
    _seed(client)
    ids = [
        _post_trade(client, closing_price=110.0 if i % 2 == 0 else 90.0)["id"]
        for i in range(4)
    ]
    rows = client.get("/api/trades").json()["rows"]
    assert [r["win_streak"] for r in rows] == [1, 0, 1, 0]  # W L W L

    response = client.post(
        "/api/trades/reorder", json={"ordered_ids": [ids[3], *ids[:3]]}
    )
    assert response.status_code == 200
    rows = client.get("/api/trades").json()["rows"]
    assert [r["win_streak"] for r in rows] == [0, 1, 0, 1]  # L W L W


def test_reorder_rejects_partial_list(client):
    _seed(client)
    ids = [_post_trade(client)["id"] for _ in range(2)]
    response = client.post("/api/trades/reorder", json={"ordered_ids": ids[:1]})
    assert response.status_code == 400


# -------------------------------------------------------- list/filter/sort/page
def test_list_search_and_filters(client):
    _seed(client)
    _post_trade(client)
    _post_trade(client, symbol="XAUUSD", closing_price=90.0, note="gold loss")

    assert client.get("/api/trades", params={"search": "gold"}).json()["total"] == 1
    assert client.get("/api/trades", params={"status": "Loss"}).json()["total"] == 1
    assert client.get("/api/trades", params={"symbol": "EURUSD"}).json()["total"] == 1
    assert (
        client.get("/api/trades", params={"opening_from": "2026-09-21"}).json()["total"]
        == 0
    )
    assert (
        client.get("/api/trades", params={"rsi_condition": "Rising"}).json()["total"]
        == 0
    )


def test_list_sort_and_pagination(client):
    _seed(client)
    for closing in (110.0, 90.0, 105.0):
        _post_trade(client, closing_price=closing)
    response = client.get("/api/trades", params={"sort": "net_pnl", "order": "desc"})
    rows = response.json()["rows"]
    assert [r["net_pnl"] for r in rows] == sorted(
        (r["net_pnl"] for r in rows), reverse=True
    )

    body = client.get("/api/trades", params={"page": 2, "page_size": 2}).json()
    assert body["total"] == 3
    assert len(body["rows"]) == 1
    assert body["page"] == 2

    assert client.get("/api/trades", params={"sort": "bogus"}).status_code == 422


# ------------------------------------------------------------------- options
def test_options_crud_and_reorder(client):
    _seed(client)
    options = client.get("/api/options").json()["lists"]
    assert [o["value"] for o in options["symbol"]] == ["EURUSD", "XAUUSD"]

    assert client.post("/api/options/symbol", json={"value": "EURUSD"}).status_code == 409
    assert client.post("/api/options/nope", json={"value": "X"}).status_code == 404

    added = client.post("/api/options/symbol", json={"value": "AUDCAD"}).json()
    response = client.patch(f"/api/options/symbol/{added['id']}", json={"value": "AUDUSD"})
    assert response.json()["value"] == "AUDUSD"

    options = client.get("/api/options").json()["lists"]["symbol"]
    ids = [o["id"] for o in options][::-1]
    assert client.post("/api/options/symbol/reorder", json={"ordered_ids": ids}).status_code == 200
    values = [o["value"] for o in client.get("/api/options").json()["lists"]["symbol"]]
    assert values == ["AUDUSD", "XAUUSD", "EURUSD"]

    assert client.delete(f"/api/options/symbol/{added['id']}").status_code == 204


def test_new_option_unlocks_trade_creation(client):
    _seed(client)
    assert client.post("/api/trades", json=dict(TRADE, symbol="AUDCAD")).status_code == 422
    client.post("/api/options/symbol", json={"value": "AUDCAD"})
    assert client.post("/api/trades", json=dict(TRADE, symbol="AUDCAD")).status_code == 201


# ------------------------------------------------------------------ analytics
def test_analytics_updates_on_change(client):
    _seed(client)
    _post_trade(client)  # win
    payload = client.get("/api/analytics").json()
    assert payload["overview"]["wins"] == 1
    assert payload["overview"]["win_rate"] == 1.0

    body = _post_trade(client, closing_price=90.0)  # loss
    payload = client.get("/api/analytics").json()
    assert payload["overview"]["wins"] == 1
    assert payload["overview"]["losses"] == 1

    client.delete(f"/api/trades/{body['id']}")
    payload = client.get("/api/analytics").json()
    assert payload["overview"]["losses"] == 0


def test_analytics_respects_filters(client):
    _seed(client)
    _post_trade(client)
    _post_trade(client, symbol="XAUUSD", closing_price=90.0)
    payload = client.get("/api/analytics", params={"symbol": "XAUUSD"}).json()
    assert payload["overview"]["total_trades"] == 1
    assert payload["overview"]["losses"] == 1


def test_group_rows_include_empty_groups(client):
    _seed(client)
    _post_trade(client)
    payload = client.get("/api/analytics").json()
    values = {row["group"]: row for row in payload["by_symbol"]}
    assert values["XAUUSD"]["total_trades"] == 0
    assert values["XAUUSD"]["win_rate"] is None


# -------------------------------------------------------------- import/export
def test_export_journal_and_backup_round_trip(client):
    import io

    from openpyxl import load_workbook

    _seed(client)
    _post_trade(client)
    journal = client.get("/api/io/export/journal")
    assert journal.status_code == 200
    ws = load_workbook(io.BytesIO(journal.content))["Journal"]
    assert ws["B1"].value == "Symbol"
    assert ws["B2"].value == "EURUSD"
    assert str(ws["AC2"].value).startswith("=IF(OR(Z2")  # live formula

    backup = client.get("/api/io/export/backup")
    assert backup.status_code == 200

    response = client.post(
        "/api/io/import?mode=replace",
        files={"file": ("backup.xlsx", backup.content)},
    )
    assert response.status_code == 200, response.text
    assert response.json()["trades_imported"] == 1
    rows = client.get("/api/trades").json()["rows"]
    assert rows[0]["symbol"] == "EURUSD"
    assert abs(rows[0]["net_pnl"] - 98.75) < 1e-9


def test_import_real_workbook_via_api(client):
    import pytest

    from conftest import REAL_WORKBOOK

    if not REAL_WORKBOOK.exists():
        pytest.skip("source workbook not available")
    with open(REAL_WORKBOOK, "rb") as fh:
        response = client.post(
            "/api/io/import?mode=replace", files={"file": ("wb.xlsx", fh)}
        )
    assert response.status_code == 200
    assert response.json()["trades_imported"] == 11
    payload = client.get("/api/analytics").json()
    assert payload["overview"]["total_trades"] == 11
    assert payload["overview"]["open_trades"] == 2


def test_import_rejects_non_xlsx(client):
    response = client.post("/api/io/import", files={"file": ("data.csv", b"a,b\n1,2")})
    assert response.status_code == 422

