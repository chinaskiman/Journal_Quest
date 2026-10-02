"""Excel importer.

Rules (per spec section 6):
  * Existing values are imported **verbatim** - no normalization, renaming,
    or recalculation.  Symbols like ``XCUUSD``, setup type ``welcome`` and
    mixed-case RSI conditions are preserved exactly.
  * Only *input* columns are read; calculated columns are always derived on
    read, never persisted.  The workbook's hardcoded ``Net PnL`` values are
    stored in the read-only ``source_net_pnl`` audit column.
  * The Analytics sheet is skipped.
  * Rows are included only when ``Symbol`` (column B) is non-empty.
"""
from __future__ import annotations

import datetime as dt
from dataclasses import dataclass, field
from pathlib import Path

from openpyxl import load_workbook
from sqlalchemy.orm import Session

from .config import EXCEL_PATH
from .models import AppState, OptionValue, Trade
from .serializers import TRADE_INPUT_FIELDS

JOURNAL_SHEET = "Journal"
OPTIONS_SHEET = "Options"

#: Excel header row -> trade field (34-column layout).
INPUT_FIELD_BY_HEADER = {
    "Symbol": "symbol",
    "Opening Day": "opening_day",
    "Opening Time": "opening_time",
    "Type": "setup_type",
    "Direction": "direction",
    "Volume": "volume",
    "Entry Price": "entry_price",
    "SL Price": "sl_price",
    "TP Price": "tp_price",
    "Risk": "risk",
    "4H Bias": "bias_4h",
    "1H Bias": "bias_1h",
    "Story Behind Candles": "story_behind_candles",
    "Opening RSI Condition": "opening_rsi_condition",
    "ATR": "atr",
    "Mental State": "mental_state",
    "Why I Entered": "why_i_entered",
    "Closing Day": "closing_day",
    "Closing Time": "closing_time",
    "Closing Price": "closing_price",
    "Closing RSI": "closing_rsi",
    "Why I Closed": "why_i_closed",
    "Commission": "commission",
    "Swap": "swap",
    "Note": "note",
}

NUMERIC_FIELDS = {
    "volume", "entry_price", "sl_price", "tp_price", "risk", "atr",
    "closing_price", "closing_rsi", "commission", "swap",
}
DATE_FIELDS = {"opening_day", "closing_day"}
TIME_FIELDS = {"opening_time", "closing_time"}

# Options sheet: header cell (row 2) -> list key
OPTIONS_HEADER_TO_KEY = {
    "Symbols": "symbol",
    "Types": "setup_type",
    "Bias": "bias",
    "RSI Condition": "rsi_condition",
    "Mental State": "mental_state",
}
OPTIONS_COLUMNS = ["A", "B", "C", "D", "E"]


@dataclass
class ImportReport:
    trades_imported: int = 0
    options_imported: int = 0
    mode: str = "replace"
    workbook: str = ""
    skipped_rows: int = 0
    notes: list[str] = field(default_factory=list)



# ---------------------------------------------------------------------------
# value converters (tolerant readers, verbatim writers)
# ---------------------------------------------------------------------------
def _to_float(value) -> float | None:
    if value is None or isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        return float(value)
    if isinstance(value, str):
        text = value.strip()
        if not text:
            return None
        try:
            return float(text)
        except ValueError:
            return None
    return None


def _to_date(value) -> dt.date | None:
    if isinstance(value, dt.datetime):
        return value.date()
    if isinstance(value, dt.date):
        return value
    if isinstance(value, str):
        text = value.strip()
        for fmt in ("%Y-%m-%d", "%m/%d/%Y", "%d/%m/%Y"):
            try:
                return dt.datetime.strptime(text, fmt).date()
            except ValueError:
                continue
    return None


def _to_time(value) -> dt.time | None:
    if isinstance(value, dt.time):
        return value
    if isinstance(value, dt.datetime):
        return value.time()
    if isinstance(value, dt.timedelta):
        return (dt.datetime.min + value).time()
    if isinstance(value, str):
        text = value.strip()
        for fmt in ("%H:%M:%S", "%H:%M"):
            try:
                return dt.datetime.strptime(text, fmt).time()
            except ValueError:
                continue
    return None


def _to_text(value) -> str | None:
    """Preserve text verbatim (no trimming!) - only '' maps to None."""
    if value is None:
        return None
    if isinstance(value, str):
        return value if value != "" else None
    return str(value)


def _convert(field: str, value):
    if field in NUMERIC_FIELDS:
        return _to_float(value)
    if field in DATE_FIELDS:
        return _to_date(value)
    if field in TIME_FIELDS:
        return _to_time(value)
    return _to_text(value)


# ---------------------------------------------------------------------------
# main entry points
# ---------------------------------------------------------------------------
def import_workbook(path: str | Path, db: Session, mode: str = "replace") -> ImportReport:
    path = Path(path)
    if not path.exists():
        raise FileNotFoundError(f"Workbook not found: {path}")

    report = ImportReport(mode=mode, workbook=str(path))
    book = load_workbook(path, data_only=True, read_only=True)

    try:
        if OPTIONS_SHEET in book.sheetnames:
            report.options_imported = _import_options(book[OPTIONS_SHEET], db, mode)
        if JOURNAL_SHEET in book.sheetnames:
            report.trades_imported = _import_journal(book[JOURNAL_SHEET], db, mode)
        else:
            raise ValueError("Workbook has no 'Journal' sheet")
    finally:
        book.close()

    db.commit()
    return report


def _import_journal(sheet, db: Session, mode: str) -> int:
    from sqlalchemy import func

    if mode == "merge":
        max_seq = db.query(func.max(Trade.seq)).scalar() or 0
    else:
        db.query(Trade).delete()
        max_seq = 0

    # Map headers -> column index (robust to inserted columns).
    header_map: dict[str, int] = {}
    for idx, cell in enumerate(next(sheet.iter_rows(min_row=1, max_row=1)), start=1):
        if cell.value is not None:
            header_map[str(cell.value).strip()] = idx
    if "Symbol" not in header_map:
        raise ValueError("Journal sheet has no 'Symbol' column")
    symbol_col = header_map["Symbol"]

    imported = 0
    n_seq = max_seq
    for row in sheet.iter_rows(min_row=2, values_only=True):
        # A row participates only if Symbol is non-empty.
        symbol_value = row[symbol_col - 1] if symbol_col - 1 < len(row) else None
        if symbol_value is None or (
            isinstance(symbol_value, str) and not symbol_value.strip()
        ):
            continue

        trade = Trade()
        for header, field_name in INPUT_FIELD_BY_HEADER.items():
            col = header_map.get(header)
            if col is None:
                continue
            raw = row[col - 1] if col - 1 < len(row) else None
            setattr(trade, field_name, _convert(field_name, raw))

        # Audit column: literal workbook Net PnL (cached value).
        net_col = header_map.get("Net PnL")
        if net_col and net_col - 1 < len(row):
            trade.source_net_pnl = _to_float(row[net_col - 1])

        n_seq += 1
        trade.seq = n_seq
        db.add(trade)
        imported += 1

    return imported


def _import_options(sheet, db: Session, mode: str) -> int:
    from openpyxl.utils import column_index_from_string

    if mode == "replace":
        db.query(OptionValue).delete()

    existing: dict[str, set[str]] = {}
    if mode == "merge":
        for key, value in db.query(OptionValue.list_key, OptionValue.value):
            existing.setdefault(key, set()).add(value)

    imported = 0
    for column_letter in OPTIONS_COLUMNS:
        col_idx = column_index_from_string(column_letter)
        header_cell = sheet.cell(row=2, column=col_idx).value
        if header_cell is None:
            continue
        list_key = OPTIONS_HEADER_TO_KEY.get(str(header_cell).strip())
        if list_key is None:
            continue

        have = existing.setdefault(list_key, set())
        position = 0
        for row_idx in range(3, sheet.max_row + 1):
            value = sheet.cell(row=row_idx, column=col_idx).value
            if value is None or str(value).strip() == "":
                break  # lists are contiguous; stop at the first blank cell
            text = str(value)
            if text in have:
                continue
            have.add(text)
            position += 1
            db.add(
                OptionValue(list_key=list_key, value=text, position=10_000 + position)
            )
            imported += 1

    return imported


def ensure_imported_on_startup(db: Session) -> ImportReport | None:
    """Import the seed workbook once, on first boot."""
    already = db.get(AppState, "excel_imported")
    if already is not None:
        return None
    if db.query(Trade).count() > 0:
        return None
    if not EXCEL_PATH.exists():
        return None
    report = import_workbook(EXCEL_PATH, db, mode="replace")
    db.add(AppState(key="excel_imported", value=str(EXCEL_PATH)))
    db.commit()
    return report